import base64
from html import escape, unescape
from io import BytesIO
import json
from pathlib import Path
import re
import time
import calendar
import os
from datetime import date, datetime, timezone
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen

import requests
import streamlit as st
from streamlit.errors import StreamlitSecretNotFoundError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

st.set_page_config(
    page_title="PLANI PRESENTACION LEVEL",
    page_icon="🎓",
    layout="wide",
)

AZUL = RGBColor(10, 42, 94)
AZUL_OSCURO = RGBColor(7, 31, 71)
AZUL_CLARO = RGBColor(191, 211, 242)
DORADO = RGBColor(212, 175, 55)
GRIS = RGBColor(51, 51, 51)
BLANCO = RGBColor(255, 255, 255)
IGNORAR_SECCIONES = {
    "referencias", "véase también", "enlaces externos", "notas", "bibliografía",
    "references", "see also", "external links", "notes", "bibliography",
    "citas", "galería", "gallery", "sources", "fuentes",
}
TABLA_PERFILES = "perfiles"
TABLA_AVISOS = "avisos"


def configuracion_supabase():
    try:
        secretos = st.secrets
        url = os.environ.get("SUPABASE_URL") or secretos.get("SUPABASE_URL", "")
        clave = os.environ.get("SUPABASE_ANON_KEY") or secretos.get(
            "SUPABASE_ANON_KEY", ""
        )
    except StreamlitSecretNotFoundError:
        url = os.environ.get("SUPABASE_URL", "")
        clave = os.environ.get("SUPABASE_ANON_KEY", "")
    return url.rstrip("/"), clave


def supabase_request(url, clave, endpoint, metodo, token=None, datos=None, params=None):
    headers = {"apikey": clave, "Content-Type": "application/json"}
    headers["Authorization"] = f"Bearer {token}" if token else f"Bearer {clave}"
    try:
        respuesta = requests.request(
            metodo,
            f"{url}{endpoint}",
            headers=headers,
            json=datos,
            params=params,
            timeout=20,
        )
    except requests.RequestException as error:
        raise RuntimeError(
            "No se pudo conectar con el servicio de cuentas. Inténtalo de nuevo."
        ) from error
    try:
        contenido = respuesta.json() if respuesta.content else None
    except requests.exceptions.JSONDecodeError as error:
        raise RuntimeError(
            "El servicio de cuentas devolvió una respuesta no válida."
        ) from error
    if not respuesta.ok:
        codigo = (contenido or {}).get("code", "")
        mensaje = (
            (contenido or {}).get("msg")
            or (contenido or {}).get("message")
            or (contenido or {}).get("error_description")
            or (contenido or {}).get("error")
            or f"Error del servicio de cuentas (HTTP {respuesta.status_code})."
        )
        if codigo in {
            "invalid_credentials",
            "invalid_grant",
            "user_not_found",
        } or "invalid login credentials" in str(mensaje).casefold():
            mensaje = "Correo electrónico o contraseña incorrectos."
        elif "email not confirmed" in str(mensaje).casefold():
            mensaje = "La cuenta no está confirmada. Contacta al administrador."
        elif respuesta.status_code == 403:
            mensaje = "No tienes autorización para realizar esta acción."
        raise RuntimeError(str(mensaje))
    return contenido


def actualizar_token(url, clave):
    sesion = st.session_state.get("sesion")
    if not sesion or not sesion.get("refresh_token"):
        raise RuntimeError("La sesión expiró. Inicia sesión de nuevo.")
    renovada = supabase_request(
        url,
        clave,
        "/auth/v1/token?grant_type=refresh_token",
        "POST",
        datos={"refresh_token": sesion["refresh_token"]},
    )
    st.session_state["sesion"] = renovada
    return renovada["access_token"]


def base_perfiles(url, clave, token, metodo="GET", datos=None, params=None):
    try:
        return supabase_request(
            url,
            clave,
            f"/rest/v1/{TABLA_PERFILES}",
            metodo,
            token,
            datos,
            params,
        )
    except RuntimeError as error:
        if "HTTP 401" not in str(error):
            raise
        nuevo_token = actualizar_token(url, clave)
        return supabase_request(
            url,
            clave,
            f"/rest/v1/{TABLA_PERFILES}",
            metodo,
            nuevo_token,
            datos,
            params,
        )


def base_avisos(url, clave, token, metodo="GET", datos=None, params=None):
    try:
        return supabase_request(
            url,
            clave,
            f"/rest/v1/{TABLA_AVISOS}",
            metodo,
            token,
            datos,
            params,
        )
    except RuntimeError as error:
        if "HTTP 401" not in str(error):
            raise
        nuevo_token = actualizar_token(url, clave)
        return supabase_request(
            url,
            clave,
            f"/rest/v1/{TABLA_AVISOS}",
            metodo,
            nuevo_token,
            datos,
            params,
        )


def perfil_actual(url, clave):
    sesion = st.session_state["sesion"]
    usuario = sesion.get("user") or {}
    filas = base_perfiles(
        url,
        clave,
        sesion["access_token"],
        params={
            "select": "id,nombre,email,rol,activo,vence",
            "id": f"eq.{usuario.get('id', '')}",
            "limit": "1",
        },
    )
    return filas[0] if filas else None


def dias_restantes(vence):
    if not vence:
        return None
    return (date.fromisoformat(vence) - date.today()).days


def sumar_un_mes(fecha):
    mes = fecha.month % 12 + 1
    anio = fecha.year + (fecha.month == 12)
    dia = min(fecha.day, calendar.monthrange(anio, mes)[1])
    return date(anio, mes, dia)


def render_login(url, clave, configurado=True):
    logo = base64.b64encode(
        Path(__file__).with_name("logo.png").read_bytes()
    ).decode("ascii")
    st.markdown(
        """
        <style>
        .stApp {
          background: radial-gradient(circle at 12% 8%, rgba(26,91,181,.72), transparent 42%),
          radial-gradient(circle at 90% 95%, rgba(212,175,55,.18), transparent 35%),
          linear-gradient(135deg, #071F47 0%, #0A2A5E 50%, #0F3F8C 100%);
        }
        [data-testid="stHeader"] { background: transparent; }
        [data-testid="stMainBlockContainer"] {
          max-width: 540px; padding-top: 2rem; padding-bottom: 3rem;
        }
        [data-testid="stForm"] {
          background: #fff; border: 0; border-radius: 18px;
          padding: 1.5rem 1.7rem; box-shadow: 0 24px 60px rgba(0,0,0,.4);
        }
        [data-testid="stForm"] label { color: #0A2A5E; font-weight: 600; }
        [data-testid="stForm"] input {
          border: 1px solid #B9C9E4; border-radius: 10px;
        }
        [data-testid="stFormSubmitButton"] button {
          width: 100% !important; background: linear-gradient(90deg, #0A2A5E, #1A5BB5);
          color: #fff; border: 0; border-radius: 10px; font-weight: 700;
        }
        .login-title {
          margin: 0; text-align: center; color: #fff;
          font-weight: 800; letter-spacing: .08em;
          font-size: clamp(1.3rem, 5vw, 1.8rem);
        }
        .login-tag {
          margin: .8rem 0 1.5rem; text-align: center; color: #BFD3F2;
        }
        .login-logo-box {
          margin-bottom: 14px; padding: 18px; text-align: center;
          background: linear-gradient(135deg, #071F47, #0F3F8C);
          border-bottom: 3px solid #D4AF37; border-radius: 16px;
        }
        .login-logo-box img { width: 130px; height: auto; max-width: 100%; }
        .login-subtitle {
          margin: .2rem 0 1.1rem; text-align: center;
          font-weight: 600; color: #0A2A5E;
        }
        .login-contact {
          margin: 1rem 0 0; text-align: center;
          color: #5A6F94; font-size: .9rem;
        }
        .login-contact a { color: #0F3F8C; font-weight: 700; }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        "<h1 class='login-title' style='color:#fff'>PLANI PRESENTACION LEVEL</h1>"
        "<p class='login-tag'>Generador de presentaciones educativas</p>",
        unsafe_allow_html=True,
    )
    with st.form("form_login", clear_on_submit=False):
        st.markdown(
            f"<div class='login-logo-box'><img src='data:image/png;base64,{logo}' alt='PLANI'></div>"
            "<p class='login-subtitle'>Inicia sesión para continuar</p>",
            unsafe_allow_html=True,
        )
        correo = st.text_input(
            "Usuario", placeholder="Correo electrónico", autocomplete="username"
        )
        contrasena = st.text_input(
            "Contraseña",
            type="password",
            autocomplete="current-password",
        )
        entrar = st.form_submit_button(
            "Aceptar", use_container_width=True
        )
        st.markdown(
            "<p class='login-contact'>Para contacto, escríbenos al "
            "<a href='tel:8094466921'>809-446-6921</a></p>",
            unsafe_allow_html=True,
        )
    if entrar:
        if not correo.strip() or not contrasena:
            st.error("Escribe tu usuario y contraseña.")
        elif not configurado:
            st.error(
                "El servicio de acceso todavía no está configurado. "
                "Contacta al administrador."
            )
        else:
            try:
                sesion = supabase_request(
                    url,
                    clave,
                    "/auth/v1/token?grant_type=password",
                    "POST",
                    datos={
                        "email": correo.strip().lower(),
                        "password": contrasena,
                    },
                )
                st.session_state["sesion"] = sesion
                st.session_state.pop("perfil", None)
                st.rerun()
            except RuntimeError as error:
                st.error(str(error))
    if configurado:
        st.stop()


def publicar_aviso(url, clave, usuarios, autor_id, tipo, destinatario, titulo, mensaje):
    titulo = titulo.strip()
    mensaje = mensaje.strip()
    if not titulo or not mensaje:
        raise RuntimeError("Escribe el título y el mensaje del aviso.")
    if len(titulo) > 100 or len(mensaje) > 1000:
        raise RuntimeError("El título admite hasta 100 caracteres y el mensaje 1,000.")
    if tipo == "individual":
        if destinatario not in {usuario["id"] for usuario in usuarios}:
            raise RuntimeError("Selecciona un usuario válido para el aviso.")
        receptores = [destinatario]
    elif tipo == "masivo":
        receptores = [None]
    else:
        raise RuntimeError("Selecciona si el aviso será individual o masivo.")

    filas = [
        {
            "destinatario_id": receptor,
            "titulo": titulo,
            "mensaje": mensaje,
            "creado_por": autor_id,
        }
        for receptor in receptores
    ]
    base_avisos(
        url,
        clave,
        st.session_state["sesion"]["access_token"],
        "POST",
        datos=filas,
        params={"select": "id"},
    )


def limpiar_avisos_vencidos(url, clave):
    base_avisos(
        url,
        clave,
        st.session_state["sesion"]["access_token"],
        "DELETE",
        params={"vence_en": f"lt.{datetime.now(timezone.utc).isoformat()}"},
    )


@st.fragment(run_every="5s")
def mostrar_avisos_activos(url, clave, usuario_id, es_admin):
    parametros = {
        "select": "id,destinatario_id,titulo,mensaje,creado_en,vence_en",
        "vence_en": f"gt.{datetime.now(timezone.utc).isoformat()}",
        "order": "creado_en.desc",
        "limit": "20",
    }
    if not es_admin:
        parametros["or"] = f"(destinatario_id.is.null,destinatario_id.eq.{usuario_id})"
    try:
        avisos = base_avisos(
            url,
            clave,
            st.session_state["sesion"]["access_token"],
            params=parametros,
        )
    except RuntimeError as error:
        st.error(f"No se pudieron consultar los avisos: {error}")
        return
    for aviso in avisos:
        st.info(f"**{aviso['titulo']}**\n\n{aviso['mensaje']}")


def guardar_perfil(url, clave, perfil, usuarios, rol, activo, vence):
    if rol not in {"admin", "usuario"}:
        raise RuntimeError("El rol seleccionado no es válido.")
    administradores_activos = [
        usuario for usuario in usuarios
        if usuario.get("rol") == "admin" and usuario.get("activo", True)
    ]
    desactiva_admin = perfil.get("rol") == "admin" and perfil.get("activo", True) and (
        rol != "admin" or not activo
    )
    if perfil["id"] == st.session_state["sesion"]["user"]["id"] and (
        rol != "admin" or not activo
    ):
        raise RuntimeError("No puedes quitarte tu propio acceso de administrador.")
    if desactiva_admin and len(administradores_activos) <= 1:
        raise RuntimeError("Debe permanecer al menos un administrador activo.")
    token = st.session_state["sesion"]["access_token"]
    base_perfiles(
        url,
        clave,
        token,
        "PATCH",
        datos={"rol": rol, "activo": activo, "vence": vence},
        params={"id": f"eq.{perfil['id']}"},
    )


@st.cache_data(ttl=3600, show_spinner=False)
def solicitar_json(url, parametros):
    consulta = urlencode(parametros)
    solicitud = Request(
        f"{url}?{consulta}",
        headers={"User-Agent": "PLANI-Presentacion-Level/1.0 (educational app)"},
    )
    for intento in range(3):
        try:
            with urlopen(solicitud, timeout=20) as respuesta:
                return json.load(respuesta)
        except HTTPError as error:
            if error.code == 429 and intento < 2:
                espera = error.headers.get("Retry-After")
                time.sleep(min(float(espera), 8) if espera and espera.isdigit() else 2 ** (intento + 1))
                continue
            raise RuntimeError(
                f"La fuente de información respondió con un error HTTP {error.code}."
            ) from error
        except (URLError, TimeoutError, json.JSONDecodeError) as error:
            raise RuntimeError(
                "No se pudo consultar una fuente de información pública. "
                "Revisa la conexión e inténtalo de nuevo."
            ) from error


def obtener_paginas(titulos, idioma, solo_intro):
    if not titulos:
        return []
    parametros = {
        "action": "query",
        "prop": "extracts|info",
        "titles": "|".join(titulos),
        "explaintext": 1,
        "exsectionformat": "wiki",
        "inprop": "url",
        "redirects": 1,
        "format": "json",
        "origin": "*",
    }
    if solo_intro:
        parametros["exintro"] = 1
        parametros["exsentences"] = 8
    datos = solicitar_json(f"https://{idioma}.wikipedia.org/w/api.php", parametros)
    paginas = (datos.get("query") or {}).get("pages") or {}
    return sorted(
        (
            {
                "titulo": pagina.get("title") or "",
                "texto": pagina.get("extract") or "",
                "url": pagina.get("fullurl") or "",
                "idioma": idioma,
            }
            for pagina in paginas.values()
            if pagina.get("extract")
        ),
        key=lambda pagina: titulos.index(pagina["titulo"])
        if pagina["titulo"] in titulos else len(titulos),
    )


def puntos_clave(texto, maximo=4):
    limpio = re.sub(r"\[\d+\]|\[[a-z]\]", "", texto)
    frases = re.split(r"(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÑ¿¡0-9])", limpio)
    return [
        frase.strip()
        for frase in frases
        if 40 <= len(frase.strip()) <= 300
    ][:maximo]


def extraer_secciones(texto):
    secciones = []
    titulo = "Introducción"
    parrafos = []
    for linea in texto.splitlines():
        linea = linea.strip()
        encabezado = re.match(r"^(={2,})\s*(.+?)\s*\1$", linea)
        if encabezado:
            if parrafos:
                secciones.append((titulo, " ".join(parrafos)))
            titulo = encabezado.group(2).strip()
            parrafos = []
        elif linea:
            parrafos.append(linea)
    if parrafos:
        secciones.append((titulo, " ".join(parrafos)))
    return [
        seccion for seccion in secciones
        if seccion[0].casefold() not in IGNORAR_SECCIONES
    ]


def buscar_contenido(tema, numero):
    paginas = []
    idioma_principal = "es"
    errores = []
    for idioma in ("es", "en"):
        try:
            datos = solicitar_json(
                f"https://{idioma}.wikipedia.org/w/api.php",
                {
                    "action": "query",
                    "generator": "search",
                    "gsrsearch": tema,
                    "gsrnamespace": 0,
                    "gsrlimit": max(numero + 6, 10),
                    "prop": "extracts|info",
                    "exintro": 1,
                    "exsentences": 8,
                    "explaintext": 1,
                    "inprop": "url",
                    "format": "json",
                    "origin": "*",
                },
            )
        except RuntimeError as error:
            errores.append(error)
            continue
        paginas = sorted(
            (
                {
                    "titulo": pagina.get("title") or "",
                    "texto": pagina.get("extract") or "",
                    "url": pagina.get("fullurl") or "",
                    "idioma": idioma,
                    "orden": pagina.get("index", 0),
                }
                for pagina in (datos.get("query") or {}).get("pages", {}).values()
                if pagina.get("extract")
            ),
            key=lambda pagina: pagina["orden"],
        )
        if paginas:
            idioma_principal = idioma
            break
    if not paginas:
        if errores:
            raise errores[-1]
        return None

    titulo_principal = paginas[0]["titulo"]
    pagina = obtener_paginas([titulo_principal], idioma_principal, False)
    texto_completo = pagina[0]["texto"] if pagina else paginas[0]["texto"]
    fuente_principal = pagina[0] if pagina else paginas[0]

    slides = []
    for titulo, texto in extraer_secciones(texto_completo):
        puntos = puntos_clave(texto)
        if puntos:
            slides.append({"titulo": titulo, "puntos": puntos, "imagen": None})
        if len(slides) >= numero:
            break

    fuentes = [{
        "nombre": f"Wikipedia ({idioma_principal}): {titulo_principal}",
        "url": fuente_principal["url"],
    }]
    usados = {titulo_principal.casefold()}
    for articulo in paginas:
        if len(slides) >= numero:
            break
        if articulo["titulo"].casefold() in usados:
            continue
        usados.add(articulo["titulo"].casefold())
        puntos = puntos_clave(articulo["texto"], 3)
        if not puntos:
            continue
        slides.append({
            "titulo": articulo["titulo"],
            "puntos": puntos,
            "imagen": None,
        })
        fuentes.append({
            "nombre": f"Wikipedia ({idioma_principal}): {articulo['titulo']}",
            "url": articulo["url"],
        })

    if len(slides) < numero and idioma_principal == "es":
        try:
            datos_en = solicitar_json(
                "https://en.wikipedia.org/w/api.php",
                {
                    "action": "query",
                    "generator": "search",
                    "gsrsearch": tema,
                    "gsrnamespace": 0,
                    "gsrlimit": numero + 5,
                    "prop": "extracts|info",
                    "exintro": 1,
                    "exsentences": 8,
                    "explaintext": 1,
                    "inprop": "url",
                    "format": "json",
                    "origin": "*",
                },
            )
            paginas_en = [
                {
                    "titulo": pagina.get("title") or "",
                    "texto": pagina.get("extract") or "",
                    "url": pagina.get("fullurl") or "",
                    "idioma": "en",
                }
                for pagina in (datos_en.get("query") or {}).get("pages", {}).values()
                if pagina.get("extract")
            ]
        except RuntimeError:
            paginas_en = []
        for articulo in paginas_en:
            if len(slides) >= numero:
                break
            if articulo["titulo"].casefold() in usados:
                continue
            usados.add(articulo["titulo"].casefold())
            puntos = puntos_clave(articulo["texto"], 3)
            if not puntos:
                continue
            slides.append({
                "titulo": articulo["titulo"],
                "puntos": puntos,
                "imagen": None,
            })
            fuentes.append({
                "nombre": f"Wikipedia (en): {articulo['titulo']}",
                "url": articulo["url"],
            })

    if not slides:
        texto_intro = paginas[0]["texto"]
        puntos = puntos_clave(texto_intro, 4)
        if puntos:
            slides.append({
                "titulo": titulo_principal,
                "puntos": puntos,
                "imagen": None,
            })
    if not slides:
        return None

    return {
        "tema": tema,
        "slides": slides[:numero],
        "fuentes": fuentes,
        "portada": None,
    }


def buscar_fuentes_extra(tema):
    fuentes = []
    avisos = []
    try:
        datos = solicitar_json(
            "https://api.openalex.org/works",
            {
                "search": tema,
                "per-page": 3,
                "select": "title,publication_year,doi",
            },
        )
        for articulo in datos.get("results", []):
            if not articulo.get("title"):
                continue
            doi = articulo.get("doi") or ""
            fuentes.append({
                "nombre": (
                    f"{articulo['title']} "
                    f"({articulo.get('publication_year') or 's.f.'})"
                ),
                "url": doi,
            })
    except RuntimeError:
        avisos.append("No se pudieron consultar artículos académicos de OpenAlex.")

    try:
        datos = solicitar_json(
            "https://openlibrary.org/search.json",
            {
                "q": tema,
                "limit": 3,
                "fields": "title,author_name,first_publish_year,key",
            },
        )
        for libro in datos.get("docs", []):
            if not libro.get("title"):
                continue
            autores = ", ".join((libro.get("author_name") or [])[:2])
            nombre = (
                f"{libro['title']} — {autores or 'Autor desconocido'} "
                f"({libro.get('first_publish_year') or 's.f.'})"
            )
            clave = libro.get("key") or ""
            fuentes.append({
                "nombre": nombre,
                "url": f"https://openlibrary.org{clave}" if clave else "",
            })
    except RuntimeError:
        avisos.append("No se pudieron consultar libros de Open Library.")
    return fuentes, avisos


def limpiar_html(texto):
    return unescape(re.sub(r"<[^>]+>", "", texto or "")).strip()


def buscar_imagenes(consulta, limite=5):
    datos = solicitar_json(
        "https://commons.wikimedia.org/w/api.php",
        {
            "action": "query",
            "generator": "search",
            "gsrnamespace": 6,
            "gsrsearch": consulta,
            "gsrlimit": limite,
            "prop": "imageinfo",
            "iiprop": "url|mime|extmetadata",
            "iiurlwidth": 1000,
            "format": "json",
            "origin": "*",
        },
    )
    paginas = (datos.get("query") or {}).get("pages") or {}
    imagenes = []
    for pagina in paginas.values():
        info = (pagina.get("imageinfo") or [{}])[0]
        url = info.get("thumburl") or info.get("url")
        if not url:
            continue
        tipo = info.get("mime")
        extension = url.split("?", 1)[0].casefold()
        if tipo not in {"image/jpeg", "image/png"} and not extension.endswith(
            (".jpg", ".jpeg", ".png")
        ):
            continue
        metadatos = info.get("extmetadata") or {}
        autor = limpiar_html(
            (metadatos.get("Artist") or {}).get("value", "")
        )[:100] or "Wikimedia Commons"
        imagenes.append({
            "url": url,
            "pagina": info.get("descriptionurl") or "",
            "autor": autor,
            "titulo": pagina.get("title") or "",
        })
    return imagenes


def descargar_imagen(imagen):
    if not imagen:
        return None
    parsed = urlparse(imagen["url"])
    if parsed.scheme != "https" or parsed.hostname not in {
        "upload.wikimedia.org",
        "thumb.wikimedia.org",
    }:
        return None
    solicitud = Request(
        imagen["url"],
        headers={"User-Agent": "PLANI-Presentacion-Level/1.0"},
    )
    try:
        with urlopen(solicitud, timeout=20) as respuesta:
            contenido = respuesta.read(12 * 1024 * 1024 + 1)
            tipo = respuesta.headers.get_content_type()
    except (HTTPError, URLError, TimeoutError, OSError):
        return None
    if len(contenido) > 12 * 1024 * 1024 or tipo not in {"image/jpeg", "image/png"}:
        return None
    return contenido


def preparar_imagenes(datos):
    candidatas = buscar_imagenes(datos["tema"], max(16, len(datos["slides"]) * 3))
    palabras_ignorar = {
        "para", "como", "desde", "sobre", "entre", "también", "este",
        "esta", "the", "and", "with", "from", "into", "file", "commons",
    }

    def palabras(texto):
        return {
            palabra for palabra in re.findall(r"[^\W_]{3,}", texto.casefold())
            if palabra not in palabras_ignorar
        }

    terminos_tema = palabras(datos["tema"])
    fuentes = set()
    usadas = set()
    imagenes_descargadas = {}
    ordenadas = sorted(
        candidatas,
        key=lambda imagen: len(terminos_tema & palabras(imagen["titulo"])),
        reverse=True,
    )
    asignaciones = [("portada", datos["tema"], None)]
    asignaciones.extend(
        ("slide", slide["titulo"], slide)
        for slide in datos["slides"]
    )
    for tipo, titulo, slide in asignaciones:
        terminos_slide = palabras(titulo)
        elegida = next(
            (
                imagen for imagen in sorted(
                    ordenadas,
                    key=lambda item: (
                        2 * len(terminos_tema & palabras(item["titulo"]))
                        + len(terminos_slide & palabras(item["titulo"]))
                    ),
                    reverse=True,
                )
                if imagen["url"] not in usadas
            ),
            None,
        )
        while elegida:
            contenido = descargar_imagen(elegida)
            usadas.add(elegida["url"])
            if contenido:
                imagen_preparada = {**elegida, "datos": contenido}
                clave = "portada" if tipo == "portada" else slide["titulo"]
                imagenes_descargadas[clave] = imagen_preparada
                if elegida["pagina"]:
                    fuentes.add((
                        f"Imagen: {elegida['autor']} (Wikimedia Commons)",
                        elegida["pagina"],
                    ))
                break
            elegida = next(
                (imagen for imagen in ordenadas if imagen["url"] not in usadas),
                None,
            )

    datos["portada"] = imagenes_descargadas.get("portada")
    for slide in datos["slides"]:
        slide["imagen"] = imagenes_descargadas.get(slide["titulo"])
    datos["fuentes"].extend(
        {"nombre": nombre, "url": url}
        for nombre, url in sorted(fuentes)
    )
    imagenes_slides = sum(
        slide["imagen"] is not None for slide in datos["slides"]
    )
    return {
        "slides": imagenes_slides,
        "total": imagenes_slides + bool(datos["portada"]),
    }


def agregar_texto(slide, texto, x, y, ancho, alto, tamano, color, negrita=False):
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(ancho), Inches(alto))
    marco = caja.text_frame
    marco.clear()
    marco.word_wrap = True
    marco.vertical_anchor = MSO_ANCHOR.MIDDLE
    parrafo = marco.paragraphs[0]
    parrafo.text = texto
    parrafo.font.name = "Calibri"
    parrafo.font.size = Pt(tamano)
    parrafo.font.bold = negrita
    parrafo.font.color.rgb = color
    return caja


def agregar_barra(slide, titulo):
    barra = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15)
    )
    barra.fill.solid()
    barra.fill.fore_color.rgb = AZUL_OSCURO
    barra.line.fill.background()
    titulo_shape = agregar_texto(
        slide, titulo, 0.55, 0.12, 12.2, 0.9,
        25 if len(titulo) > 48 else 30, BLANCO, True,
    )
    titulo_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    acento = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.15), Inches(13.333), Inches(0.07)
    )
    acento.fill.solid()
    acento.fill.fore_color.rgb = DORADO
    acento.line.fill.background()


def agregar_pie(slide, numero):
    agregar_texto(
        slide, f"PLANI PRESENTACION LEVEL  |  {numero}",
        9.3, 7.12, 3.5, 0.2, 9, AZUL,
    )


def agregar_imagen(slide, imagen, x, y, ancho_max, alto_max):
    archivo = BytesIO(imagen["datos"])
    from PIL import Image

    with Image.open(archivo) as foto:
        proporcion = foto.width / foto.height
    ancho = min(ancho_max, alto_max * proporcion)
    alto = min(alto_max, ancho_max / proporcion)
    archivo.seek(0)
    slide.shapes.add_picture(
        archivo,
        Inches(x + (ancho_max - ancho) / 2),
        Inches(y + (alto_max - alto) / 2),
        width=Inches(ancho),
        height=Inches(alto),
    )
    return ancho, alto


def ajustar_tamano_puntos(puntos):
    total = sum(len(punto) for punto in puntos)
    if total > 900:
        return 16
    if total > 700:
        return 17
    if total > 500:
        return 18
    return 20


def crear_presentacion(tema, slides, fuentes, portada, incluir_videos):
    presentacion = Presentation()
    presentacion.slide_width = Inches(13.333)
    presentacion.slide_height = Inches(7.5)
    presentacion.core_properties.title = tema
    presentacion.core_properties.author = "PLANI PRESENTACION LEVEL"
    disposicion_vacia = presentacion.slide_layouts[6]

    portada_slide = presentacion.slides.add_slide(disposicion_vacia)
    portada_slide.background.fill.solid()
    portada_slide.background.fill.fore_color.rgb = AZUL_OSCURO
    acento = portada_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.2), Inches(0.09), Inches(5.1)
    )
    acento.fill.solid()
    acento.fill.fore_color.rgb = DORADO
    acento.line.fill.background()
    ancho_titulo = 6.8 if portada else 11.4
    agregar_texto(
        portada_slide, tema.strip().capitalize(), 1.05, 1.7,
        ancho_titulo, 2.8, 44 if len(tema) < 34 else 34, BLANCO, True,
    )
    agregar_texto(
        portada_slide, "Presentación educativa basada en fuentes públicas",
        1.08, 4.7, ancho_titulo, 0.65, 19, AZUL_CLARO,
    )
    agregar_texto(
        portada_slide, "PLANI PRESENTACION LEVEL", 1.08, 6.75,
        7, 0.3, 12, DORADO, True,
    )
    if portada:
        agregar_imagen(portada_slide, portada, 8.0, 1.2, 4.6, 4.9)
        agregar_texto(
            portada_slide, f"Imagen: {portada['autor']}",
            8.0, 6.15, 4.6, 0.3, 10, AZUL_CLARO,
        )

    for indice, contenido in enumerate(slides, start=1):
        slide = presentacion.slides.add_slide(disposicion_vacia)
        agregar_barra(slide, contenido["titulo"])
        con_imagen = contenido["imagen"] is not None
        ancho_texto = 7.3 if con_imagen else 12.2
        alto = 5.35 / max(len(contenido["puntos"]), 1)
        tamano = ajustar_tamano_puntos(contenido["puntos"])
        for punto_indice, punto in enumerate(contenido["puntos"]):
            caja = slide.shapes.add_textbox(
                Inches(0.7),
                Inches(1.55 + punto_indice * alto),
                Inches(ancho_texto),
                Inches(alto),
            )
            marco = caja.text_frame
            marco.clear()
            marco.word_wrap = True
            parrafo = marco.paragraphs[0]
            parrafo.level = 0
            parrafo.font.name = "Calibri"
            parrafo.font.size = Pt(tamano)
            parrafo.font.color.rgb = GRIS
            parrafo.space_after = Pt(9)
            parrafo.text = f"• {punto}"
        if con_imagen:
            imagen = contenido["imagen"]
            _, alto_imagen = agregar_imagen(slide, imagen, 8.35, 1.55, 4.35, 4.8)
            agregar_texto(
                slide, f"Imagen: {imagen['autor']}",
                8.25, min(6.45, 1.6 + alto_imagen), 4.6, 0.35, 9, AZUL,
            )
        agregar_pie(slide, indice)

    referencias = presentacion.slides.add_slide(disposicion_vacia)
    agregar_barra(referencias, "Fuentes y referencias")
    fuentes_vistas = set()
    lineas = []
    for fuente in fuentes:
        clave = (fuente["nombre"], fuente["url"])
        if clave not in fuentes_vistas:
            fuentes_vistas.add(clave)
            lineas.append(fuente)
        if len(lineas) >= 14:
            break
    for indice, fuente in enumerate(lineas):
        y = 1.5 + indice * min(0.39, 5.45 / max(len(lineas), 1))
        caja = referencias.shapes.add_textbox(
            Inches(0.7), Inches(y), Inches(12), Inches(0.42)
        )
        caja.text_frame.clear()
        parrafo = caja.text_frame.paragraphs[0]
        parrafo.text = fuente["nombre"] + (f" — {fuente['url']}" if fuente["url"] else "")
        parrafo.font.name = "Calibri"
        parrafo.font.size = Pt(12)
        parrafo.font.color.rgb = AZUL
        if fuente["url"]:
            parrafo.runs[0].hyperlink.address = fuente["url"]

    if incluir_videos:
        videos = presentacion.slides.add_slide(disposicion_vacia)
        agregar_barra(videos, "Videos recomendados en YouTube")
        consultas = [
            ("Explicación del tema", tema),
            ("Clases y tutoriales", f"{tema} clase educativa"),
        ]
        for indice, (etiqueta, consulta) in enumerate(consultas):
            url = "https://www.youtube.com/results?" + urlencode(
                {"search_query": consulta}
            )
            caja = videos.shapes.add_textbox(
                Inches(0.8), Inches(1.8 + indice * 0.9), Inches(11.7), Inches(0.55)
            )
            parrafo = caja.text_frame.paragraphs[0]
            parrafo.text = f"{etiqueta}: {url}"
            parrafo.font.name = "Calibri"
            parrafo.font.size = Pt(18)
            parrafo.font.color.rgb = AZUL
            parrafo.runs[0].hyperlink.address = url

    salida = BytesIO()
    presentacion.save(salida)
    return salida.getvalue()


url_supabase, clave_supabase = configuracion_supabase()
st.markdown(
    """
    <style>
    .stApp { background: #EEF3FA; color: #0A2A5E; }
    [data-testid="stHeader"] { background: rgba(238, 243, 250, .92); }
    [data-testid="stFooter"], footer { display: none !important; }
    .plani-banner {
        background: linear-gradient(120deg, #071F47 0%, #0A2A5E 48%, #1A5BB5 100%);
        color: white; padding: 1.5rem 2rem; border-radius: 20px;
        border-bottom: 5px solid #D4AF37; margin: 0 auto 1.3rem;
        box-shadow: 0 10px 28px rgba(10, 42, 94, .23);
    }
    .plani-banner h1 { color: white; margin: 0; font-size: 1.8rem; }
    .plani-banner p { color: #BFD3F2; margin: .45rem 0 0; }
    .plani-user {
        margin: .25rem 0 1rem; padding: 1rem .75rem; text-align: center;
        border: 1px solid rgba(10, 42, 94, .12); border-radius: 14px;
        background: white;
    }
    .plani-user-avatar {
        width: 52px; height: 52px; margin: 0 auto .55rem; border-radius: 50%;
        display: grid; place-items: center; background: #D4AF37; color: #071F47;
        font-size: 1.25rem; font-weight: 800;
    }
    .plani-user-name { color: #0A2A5E; font-weight: 700; overflow-wrap: anywhere; }
    .plani-user-status { margin-top: .25rem; color: #16834A; font-size: .85rem; }
    .plani-user-status::before {
        content: ""; display: inline-block; width: 8px; height: 8px; margin-right: .4rem;
        border-radius: 50%; background: #20A464; vertical-align: middle;
    }
    [data-testid="stForm"] {
        background: white; padding: 1.5rem 1.7rem; border-radius: 18px;
        border: 1px solid #D9E4F5; box-shadow: 0 8px 24px rgba(10, 42, 94, .1);
    }
    [data-testid="stFormSubmitButton"] button, [data-testid="stDownloadButton"] button {
        background: linear-gradient(90deg, #0A2A5E, #1A5BB5);
        color: white; border: 0; font-weight: 700; border-radius: 10px;
    }
    [data-testid="stSlider"] [role="slider"] { background-color: #0F3F8C; }
    </style>
    """,
    unsafe_allow_html=True,
)

if not url_supabase or not clave_supabase:
    render_login("", "", configurado=False)
    st.stop()

if not st.session_state.get("sesion"):
    render_login(url_supabase, clave_supabase)

try:
    perfil = perfil_actual(url_supabase, clave_supabase)
except RuntimeError as error:
    st.error(str(error))
    if st.button("Volver al inicio de sesión"):
        st.session_state.pop("sesion", None)
        st.rerun()
    st.stop()

if not perfil:
    st.error(
        "Tu cuenta todavía no tiene un perfil autorizado. "
        "Contacta al administrador para que la habilite."
    )
    if st.button("Cerrar sesión"):
        st.session_state.pop("sesion", None)
        st.rerun()
    st.stop()

if not perfil.get("activo", True):
    st.error("Tu cuenta está desactivada. Contacta al administrador.")
    if st.button("Cerrar sesión"):
        st.session_state.pop("sesion", None)
        st.rerun()
    st.stop()

es_admin = perfil.get("rol") == "admin"
dias = dias_restantes(perfil.get("vence"))
vencido = dias is not None and dias <= 0
if vencido and not es_admin:
    st.error(
        "Tu acceso venció. El administrador debe autorizar la renovación "
        "mensual para que puedas volver a entrar."
    )
    if st.button("Cerrar sesión"):
        st.session_state.pop("sesion", None)
        st.session_state.pop("presentacion", None)
        st.rerun()
    st.stop()

st.markdown(
    """
    <div class="plani-banner">
      <h1>🎓 PLANI PRESENTACION LEVEL</h1>
      <p>Generador de presentaciones educativas con información real, imágenes y referencias</p>
    </div>
    """,
    unsafe_allow_html=True,
)

mostrar_avisos_activos(
    url_supabase,
    clave_supabase,
    perfil["id"],
    es_admin,
)

with st.sidebar:
    nombre_usuario = str(perfil.get("nombre") or perfil.get("email") or "Usuario")
    iniciales = "".join(
        parte[0] for parte in nombre_usuario.strip().split()[:2] if parte
    ).upper() or "U"
    st.markdown(
        "<div class='plani-user'>"
        f"<div class='plani-user-avatar'>{escape(iniciales)}</div>"
        f"<div class='plani-user-name'>{escape(nombre_usuario)}</div>"
        "<div class='plani-user-status'>En línea</div>"
        "</div>",
        unsafe_allow_html=True,
    )
    st.caption("Administrador" if es_admin else "Usuario")
    if not perfil.get("vence"):
        st.success("Acceso sin vencimiento")
    elif vencido:
        st.error("Acceso vencido; puedes administrar renovaciones.")
    else:
        st.info(
            f"Quedan {dias} día(s). Vence el {perfil['vence']}."
        )
    if st.button("Cerrar sesión", use_container_width=True):
        st.session_state.pop("sesion", None)
        st.session_state.pop("presentacion", None)
        st.rerun()
    opciones = ["Generar presentación"]
    if es_admin:
        opciones.append("Administrar usuarios")
        opciones.append("Avisos")
    seccion = st.radio("Panel", opciones, label_visibility="collapsed")

if seccion == "Avisos":
    st.subheader("📣 Enviar avisos")
    st.write(
        "Los avisos individuales o masivos se muestran a sus destinatarios "
        "durante un minuto y luego desaparecen."
    )
    try:
        limpiar_avisos_vencidos(url_supabase, clave_supabase)
        usuarios_aviso = base_perfiles(
            url_supabase,
            clave_supabase,
            st.session_state["sesion"]["access_token"],
            params={
                "select": "id,nombre,email,activo",
                "order": "email.asc",
            },
        )
    except RuntimeError as error:
        st.error(f"No se pudieron cargar los usuarios: {error}")
        st.stop()

    if not usuarios_aviso:
        st.info("No hay usuarios a quienes enviar avisos.")
        st.stop()

    usuario_por_id = {usuario["id"]: usuario for usuario in usuarios_aviso}
    with st.form("form_aviso"):
        tipo_aviso = st.radio(
            "Destinatarios",
            ["Individual", "Masivo (todos los usuarios)"],
            horizontal=True,
        )
        destinatario_aviso = None
        if tipo_aviso == "Individual":
            destinatario_aviso = st.selectbox(
                "Usuario",
                list(usuario_por_id),
                format_func=lambda uid: (
                    f"{usuario_por_id[uid].get('nombre') or usuario_por_id[uid].get('email')} "
                    f"— {usuario_por_id[uid].get('email')}"
                ),
            )
        titulo_aviso = st.text_input("Título", max_chars=100)
        mensaje_aviso = st.text_area("Mensaje", max_chars=1000, height=140)
        enviar_aviso = st.form_submit_button("Enviar aviso", use_container_width=True)
    if enviar_aviso:
        try:
            publicar_aviso(
                url_supabase,
                clave_supabase,
                usuarios_aviso,
                perfil["id"],
                "individual" if tipo_aviso == "Individual" else "masivo",
                destinatario_aviso,
                titulo_aviso,
                mensaje_aviso,
            )
            st.success("Aviso enviado. Dejará de mostrarse después de un minuto.")
        except RuntimeError as error:
            st.error(str(error))

elif seccion == "Administrar usuarios":
    st.subheader("👥 Panel de administrador")
    st.write(
        "Consulta el contador de cada cuenta, ajusta vencimientos y autoriza "
        "renovaciones de un mes."
    )
    try:
        usuarios = base_perfiles(
            url_supabase,
            clave_supabase,
            st.session_state["sesion"]["access_token"],
            params={
                "select": "id,nombre,email,rol,activo,vence",
                "order": "email.asc",
            },
        )
    except RuntimeError as error:
        st.error(str(error))
        st.stop()

    if not usuarios:
        st.info("No hay perfiles creados todavía.")
        st.stop()

    def estado_acceso(usuario):
        restante = dias_restantes(usuario.get("vence"))
        if not usuario.get("activo", True):
            return "Desactivado", "—"
        if restante is None:
            return "Sin vencimiento", "—"
        if restante <= 0:
            return "Vencido", "0 días"
        return "Activo", f"{restante} días"

    estados = [estado_acceso(usuario) for usuario in usuarios]
    metricas = st.columns(4)
    metricas[0].metric("Usuarios", len(usuarios))
    metricas[1].metric(
        "Con acceso", sum(estado == "Activo" or estado == "Sin vencimiento" for estado, _ in estados)
    )
    metricas[2].metric("Vencidos", sum(estado == "Vencido" for estado, _ in estados))
    metricas[3].metric("Desactivados", sum(estado == "Desactivado" for estado, _ in estados))

    st.dataframe(
        [
            {
                "Nombre": usuario.get("nombre") or "",
                "Correo": usuario.get("email") or "",
                "Rol": "Administrador" if usuario.get("rol") == "admin" else "Usuario",
                "Estado": estado,
                "Tiempo restante": restante,
                "Vence": usuario.get("vence") or "Sin vencimiento",
            }
            for usuario, (estado, restante) in zip(usuarios, estados)
        ],
        hide_index=True,
        use_container_width=True,
    )

    por_id = {usuario["id"]: usuario for usuario in usuarios}
    seleccionado_id = st.selectbox(
        "Seleccionar cuenta para administrar",
        list(por_id),
        format_func=lambda uid: (
            f"{por_id[uid].get('nombre') or por_id[uid].get('email')} "
            f"— {por_id[uid].get('email')}"
        ),
    )
    seleccionado = por_id[seleccionado_id]
    st.markdown(f"#### {seleccionado.get('email', '')}")
    vence_actual = seleccionado.get("vence")
    fecha_actual = date.fromisoformat(vence_actual) if vence_actual else date.today()
    with st.form("form_editar_acceso"):
        nuevo_rol = st.selectbox(
            "Rol",
            ["usuario", "admin"],
            index=0 if seleccionado.get("rol") != "admin" else 1,
            format_func=lambda valor: "Administrador" if valor == "admin" else "Usuario",
        )
        activa = st.checkbox(
            "Cuenta activa",
            value=seleccionado.get("activo", True),
        )
        sin_vence = st.checkbox(
            "Sin fecha de vencimiento",
            value=vence_actual is None,
        )
        nueva_fecha = st.date_input(
            "Fecha de vencimiento",
            value=fecha_actual,
            disabled=sin_vence,
        )
        guardar = st.form_submit_button("Guardar cambios")
    if guardar:
        try:
            guardar_perfil(
                url_supabase,
                clave_supabase,
                seleccionado,
                usuarios,
                nuevo_rol,
                activa,
                None if sin_vence else nueva_fecha.isoformat(),
            )
            st.success("Cambios de acceso guardados.")
            st.rerun()
        except RuntimeError as error:
            st.error(str(error))

    st.caption(
        "La renovación se realiza solo al pulsar el botón de autorización. "
        "Si el vencimiento sigue vigente, agrega un mes a esa fecha; si ya venció, "
        "cuenta un mes desde hoy."
    )
    if st.button("✅ Autorizar renovación de un mes"):
        base = date.fromisoformat(vence_actual) if vence_actual else date.today()
        nueva_vence = sumar_un_mes(max(base, date.today()))
        try:
            base_perfiles(
                url_supabase,
                clave_supabase,
                st.session_state["sesion"]["access_token"],
                "PATCH",
                datos={"vence": nueva_vence.isoformat(), "activo": True},
                params={"id": f"eq.{seleccionado_id}"},
            )
            st.success(
                f"Renovación autorizada. El acceso ahora vence el {nueva_vence.isoformat()}."
            )
            st.rerun()
        except RuntimeError as error:
            st.error(str(error))
else:
    st.markdown(
        "Escribe un tema y crea un PowerPoint con información de Wikipedia, "
        "imágenes de Wikimedia Commons, referencias y enlaces de YouTube."
    )
    if vencido:
        st.warning(
            "Tu cuenta de administrador venció: puedes renovar usuarios, "
            "pero no generar presentaciones hasta renovar tu propio acceso."
        )
    with st.form("form_presentacion"):
        tema = st.text_input(
            "Tema, asignatura o concepto",
            "La fotosíntesis",
            max_chars=120,
            disabled=vencido,
        )
        numero = st.slider(
            "Diapositivas de contenido", 3, 10, 6, disabled=vencido
        )
        incluir_videos = st.checkbox(
            "Incluir enlaces de YouTube", value=True, disabled=vencido
        )
        generar = st.form_submit_button(
            "🚀 Generar presentación", disabled=vencido
        )

    if generar:
        if not tema.strip():
            st.session_state.pop("presentacion", None)
            st.error("Escribe un tema para continuar.")
        else:
            st.session_state.pop("presentacion", None)
            try:
                with st.spinner("Investigando el tema y buscando imágenes relacionadas…"):
                    datos = buscar_contenido(tema.strip(), numero)
                    if not datos:
                        st.error(
                            "No se encontró información suficiente. "
                            "Prueba con otro tema o una redacción diferente."
                        )
                    else:
                        fuentes_extra, avisos_fuentes = buscar_fuentes_extra(
                            datos["tema"]
                        )
                        datos["fuentes"].extend(fuentes_extra)
                        imagenes = preparar_imagenes(datos)
                        archivo = crear_presentacion(
                            datos["tema"],
                            datos["slides"],
                            datos["fuentes"],
                            datos["portada"],
                            incluir_videos,
                        )
                        nombre = re.sub(
                            r"[^\w-]+", "_", tema.strip(), flags=re.UNICODE
                        ).strip("_")[:60] or "presentacion"
                        st.session_state["presentacion"] = {
                            "datos": archivo,
                            "nombre": f"{nombre}.pptx",
                            "diapositivas": len(datos["slides"]),
                            "imagenes": imagenes["total"],
                            "imagenes_slides": imagenes["slides"],
                            "avisos_fuentes": avisos_fuentes,
                        }
            except RuntimeError as error:
                st.error(str(error))
            except (ValueError, OSError) as error:
                st.error(f"No se pudo armar el archivo PowerPoint: {error}")

    presentacion = st.session_state.get("presentacion")
    if presentacion:
        st.success(
            f"Presentación lista: {presentacion['diapositivas']} diapositivas de contenido "
            f"y {presentacion['imagenes']} imágenes relacionadas con el tema."
        )
        if presentacion["imagenes_slides"] < presentacion["diapositivas"]:
            st.info(
                "No se encontraron imágenes públicas para todas las diapositivas; "
                "el resto conserva el diseño azul y su contenido."
            )
        for aviso in presentacion["avisos_fuentes"]:
            st.warning(aviso)
        st.download_button(
            label="📥 Descargar presentación (.pptx)",
            data=presentacion["datos"],
            file_name=presentacion["nombre"],
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
